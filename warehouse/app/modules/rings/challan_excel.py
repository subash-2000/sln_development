"""SLN delivery challan Excel, filled into the original Buyer1 sheet layout.

Buyer1 and Buyer2 use the same sheet. Only the buyer name and address change.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from app.models.company import Company
from app.models.rings import RingsDocument
from app.modules.rings.calc import packing_detail_text
from app.modules.rings.pdf_common import amount_in_words, sanitize_po

TEMPLATE = Path(__file__).resolve().parent / "templates" / "sln_challan.xlsx"
DATA_FIRST = 15
DATA_LAST = 32


def _sheet_title(document: RingsDocument, used: set[str]) -> str:
    raw = (document.sln_challan_number or document.document_number or "Challan").strip()
    cleaned = "".join(ch for ch in raw if ch not in "\\/*?:[]")[:28] or "Challan"
    title = cleaned
    suffix = 2
    while title in used or title == "_template":
        title = f"{cleaned[:24]}-{suffix}"
        suffix += 1
    used.add(title)
    return title


def _clean(value: str | None) -> str:
    return " ".join((value or "").replace("\r", " ").split())


def _party_address(party) -> str:
    if party is None:
        return ""
    return _clean(getattr(party, "billing_address", None) or getattr(party, "shipping_address", None))


def _excel_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return None


def _lines(*parts: str) -> str:
    return "\n".join(part for part in parts if part)


def _plain_code(value) -> str | None:
    """Keep codes such as HSN as text, without a trailing .0."""
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        number = Decimal(raw)
    except Exception:
        return raw
    if number == number.to_integral():
        return str(int(number))
    return format(number, "f")


def _fit(cell, value, *, as_text: bool = False) -> None:
    """Show the full value on one line inside the existing row."""
    if isinstance(value, str):
        value = " ".join(value.replace("\r", " ").replace("\n", " ").split())
    cell.value = value
    if as_text:
        cell.number_format = "@"
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False, shrink_to_fit=True)


def _fill(ws: Worksheet, document: RingsDocument, company: Company | None) -> None:
    ws.sheet_view.showGridLines = True
    ws.sheet_view.zoomScale = 70
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = "B2:N45"
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.view = "normal"

    vendor = document.vendor
    customer = document.customer
    supplier_name = _clean(vendor.name if vendor else "")
    buyer_name = _clean(customer.name if customer else "")

    if company is not None:
        ws["B5"] = _lines(
            _clean(company.legal_name or company.name),
            _clean(company.address),
            f"GSTIN/UIN : {company.gstin}" if company.gstin else "",
        )

    ws["K4"] = document.sln_challan_number or ""
    ws["K4"].number_format = "@"
    ws["M4"] = _excel_date(document.sln_challan_date or document.created_at)
    ws["M4"].number_format = "DD-MM-YYYY"

    parent = document.__dict__.get("parent_document")
    inward_no = ""
    inward_date = document.received_date
    if parent is not None:
        inward_no = getattr(parent, "document_number", "") or ""
        inward_date = getattr(parent, "received_date", None) or getattr(parent, "supplier_document_date", None) or inward_date
    ws["K5"] = inward_no or None
    ws["M5"] = _excel_date(inward_date)
    if inward_date is not None:
        ws["M5"].number_format = "DD-MM-YYYY"

    ws["B7"] = _lines(
        buyer_name,
        _party_address(customer),
        f"GSTIN/UIN : {customer.gstin}" if customer and customer.gstin else "",
        f"PAN/IT No : {customer.pan}" if customer and customer.pan else "",
    )
    # Transporter stays blank so it can be written by hand. Vehicle and EWB are filled.
    ws["M7"] = document.vehicle_number or None
    ws["K8"] = document.eway_bill_number or None

    ws["B9"] = _lines(
        "PRINCIPAL SUPPLIER",
        supplier_name,
        _party_address(vendor),
        f"GSTIN :  {vendor.gstin}" if vendor and vendor.gstin else "",
    )

    invoice_no = _clean(document.supplier_customer_invoice_number) or "___________"
    invoice_dt = "_____"
    if document.supplier_customer_invoice_number and document.supplier_customer_invoice_date:
        invoice_dt = document.supplier_customer_invoice_date.strftime("%d-%m-%Y")
    ws["J9"] = (
        f"THIS DELIVERY CHALLAN HAS REFERENCE TO INVOICE NO {invoice_no} DT. {invoice_dt}   "
        f"OF THE PRINCIPAL SUPPLIER, {supplier_name}"
    )

    lines = list(document.lines or [])
    extra = max(0, len(lines) - (DATA_LAST - DATA_FIRST + 1))
    if extra:
        ws.insert_rows(DATA_LAST + 1, extra)
    last_data_row = DATA_LAST + extra

    for row in range(DATA_FIRST, last_data_row + 1):
        ws.row_dimensions[row].height = 29.25
        for col in range(2, 15):
            ws.cell(row, col).value = None

    total_amt = Decimal("0")
    for index, line in enumerate(lines):
        row = DATA_FIRST + index
        packing = packing_detail_text(
            line.packing_detail,
            line.pallet_count,
            line.qty_per_pallet,
            line.box_count,
            line.qty_per_box,
        )
        hsn = _plain_code(line.product.hsn_code if line.product else "")
        qty = line.quantity or Decimal("0")
        rate = line.unit_rate or Decimal("0")
        amount = line.basic_value if line.basic_value is not None else (qty * rate)
        total_amt += Decimal(str(amount))
        challan_no = document.supplier_document_number or ""
        ws.cell(row, 2, index + 1)
        ws.cell(row, 2).alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        ws.cell(row, 3, (line.product.sku if line.product else "") or None)
        ws.cell(row, 3).alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        _fit(ws.cell(row, 5), hsn, as_text=True)
        ws.cell(row, 6, int(challan_no) if str(challan_no).isdigit() else (challan_no or None))
        ws.cell(row, 7, _excel_date(document.supplier_document_date))
        ws.cell(row, 7).number_format = "DD-MM-YYYY"
        po = sanitize_po(document.po_number)
        ws.cell(row, 8, po or None)
        ws.cell(row, 9, _excel_date(document.po_date))
        if document.po_date is not None:
            ws.cell(row, 9).number_format = "DD-MM-YYYY"
        _fit(ws.cell(row, 10), line.heat_number or None)
        ws.cell(row, 11, packing or None)
        ws.cell(row, 12, float(qty))
        ws.cell(row, 13, float(rate))
        ws.cell(row, 14, f"=L{row}*M{row}")

    max_heat_len = max([len(str(getattr(l, "heat_number", "") or "").strip()) for l in lines] + [len("Heat No.")])
    ws.column_dimensions["B"].width = 20.0
    ws.column_dimensions["E"].width = 24.0
    ws.column_dimensions["J"].width = max(48.0, float(int(max_heat_len * 1.6 + 6)))

    total_row = last_data_row + 1
    ws.cell(total_row, 12, f"=SUM(L{DATA_FIRST}:L{last_data_row})")
    ws.cell(total_row, 14, f"=SUM(N{DATA_FIRST}:N{last_data_row})")
    words_row = total_row + 1
    ws.cell(words_row, 4, amount_in_words(total_amt))

    tax_row = total_row + 3
    ws.row_dimensions[tax_row].height = 22.5
    ws.cell(tax_row - 1, 2).alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    hsn_show = _plain_code(lines[0].product.hsn_code) if lines and lines[0].product else None
    _fit(ws.cell(tax_row, 2), hsn_show, as_text=True)
    ws.cell(tax_row, 3, f"=N{total_row}")
    ws.cell(tax_row, 7, 0)
    ws.cell(tax_row, 7).number_format = "0%"
    ws.cell(tax_row, 8, 0)
    ws.cell(tax_row, 12, 0)


def _workbook(documents: list[RingsDocument], company: Company | None):
    wb = load_workbook(TEMPLATE)
    pristine = wb.active
    pristine.title = "_template"
    used: set[str] = set()
    for document in documents:
        ws = wb.copy_worksheet(pristine)
        _fill(ws, document, company)
        ws.title = _sheet_title(document, used)
    del wb["_template"]
    if wb.worksheets:
        wb.active = wb.worksheets[0]
    return wb


def generate_sln_challan_xlsx(document: RingsDocument, company: Company | None) -> bytes:
    buffer = BytesIO()
    _workbook([document], company).save(buffer)
    return buffer.getvalue()


def generate_sln_challan_register_xlsx(documents: list[RingsDocument], company: Company | None) -> bytes:
    buffer = BytesIO()
    _workbook(documents, company).save(buffer)
    return buffer.getvalue()
