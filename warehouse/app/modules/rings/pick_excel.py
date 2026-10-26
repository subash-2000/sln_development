"""Excel picking list in the '2nd Picking List' sheet layout."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.rings import RingsDocument
from app.modules.rings.pick_layout import (
    group_picks_by_buyer,
    pick_headers,
    pick_line_values,
    pick_sheet_date,
)

THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
LAVENDER_FILL = PatternFill("solid", fgColor="CCC1DA")
HEADER_FILL = PatternFill("solid", fgColor="FFFF00")
TITLE_FONT = Font(bold=True, size=18, name="Calibri", color="000000")
DATE_FONT = Font(bold=True, size=16, name="Calibri", color="000000")
HEADER_FONT = Font(bold=True, color="000000", size=10, name="Calibri")
BODY_FONT = Font(name="Calibri", size=10)


def _prepare(ws) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.oddHeader.center.text = "Picking Slip"
    ws.print_title_rows = "1:4"
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = None


def _write_block(ws, row: int, buyer_name: str, docs: list[RingsDocument], date_label: str) -> int:
    workflow = docs[0].workflow if docs else None
    headers = pick_headers(workflow)
    width = max(len(headers), 1)

    # Date row: merged, lavender background, centered (like Agrasen template)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    for col in range(1, width + 1):
        c = ws.cell(row, col)
        c.fill = LAVENDER_FILL
    date_cell = ws.cell(row, 1, date_label)
    date_cell.font = DATE_FONT
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 26
    row += 2

    # Headers: yellow background, bold, centered
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row, col, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = THIN
    ws.row_dimensions[row].height = 36
    row += 1

    # Data rows: single line, no wrapping, centered
    for doc in docs:
        for line in doc.lines or []:
            values = pick_line_values(doc, line)
            for col, value in enumerate(values, start=1):
                clean_val = value
                if isinstance(clean_val, str):
                    clean_val = " ".join(clean_val.replace("\r", " ").replace("\n", " ").split())
                cell = ws.cell(row, col, clean_val if clean_val != "" else None)
                cell.border = THIN
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
                cell.font = BODY_FONT
            ws.row_dimensions[row].height = 24
            row += 1
    if buyer_name and row > 0:
        return row + 2
    return row + 2


def generate_pick_slip_xlsx(document: RingsDocument) -> bytes:
    return generate_pick_register_xlsx([document])


def generate_pick_register_xlsx(documents: list[RingsDocument]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "2nd Picking List"
    _prepare(ws)

    headers = pick_headers(documents[0].workflow if documents else None)
    width = max(len(headers), 1)

    # Row 1: "Picking Slip" title bar with lavender fill
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    for col in range(1, width + 1):
        c = ws.cell(1, col)
        c.fill = LAVENDER_FILL
    title = ws.cell(1, 1, "Picking Slip")
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    groups = group_picks_by_buyer(documents) or [("Unassigned buyer", [])]
    row = 2
    for buyer_name, docs in groups:
        row = _write_block(ws, row, buyer_name, docs, pick_sheet_date(docs or documents))

    # Calculate max heat number length to ensure Column width never wraps
    max_heat_len = 0
    for doc in documents:
        for line in doc.lines or []:
            if line.heat_number:
                max_heat_len = max(max_heat_len, len(str(line.heat_number).strip()))
    heat_width = max(38.0, float(max_heat_len + 5))

    header_widths = {
        "TYPES": 15.0,
        "Challan NO": 16.0,
        "Invoice Number": 18.0,
        "DATE": 14.0,
        "TRANSPORTER": 20.0,
        "REC.DATE": 14.0,
        "Total Pallets": 15.0,
        "Per Pallet * Qty": 18.0,
        "Total QTY": 15.0,
        "Supplier Name": 24.0,
        "Buyer Name": 32.0,
        "Sln To Supplier Challan Number": 26.0,
        "Supplier To Customer Invoice Number": 26.0,
        "Purchase Order Number": 24.0,
        "Heat Number": heat_width,
        "Per Ring Weight": 16.0,
        "Packing Weight": 16.0,
        "Gross Weight": 16.0,
    }

    for index, header_name in enumerate(headers, start=1):
        col_letter = get_column_letter(index)
        ws.column_dimensions[col_letter].width = header_widths.get(header_name, 18.0)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
